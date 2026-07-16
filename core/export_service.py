import re
from pathlib import Path

from PyQt6.QtGui import QTextDocument
from PyQt6.QtPrintSupport import QPrinter
from docx import Document

from core.file_service import write_text_file_atomic

_BOLD_ITALIC = re.compile(r"(\*\*.+?\*\*|\*.+?\*)")
_HEADING = re.compile(r"^(#{1,6})\s+(.*)")
_BULLET = re.compile(r"^[-*]\s+(.*)")
_NUMBERED = re.compile(r"^\d+\.\s+(.*)")


def export_to_pdf(markdown_text: str, path: str):
    document = QTextDocument()
    document.setMarkdown(markdown_text)
    printer = QPrinter(QPrinter.PrinterMode.HighResolution)
    printer.setOutputFormat(QPrinter.OutputFormat.PdfFormat)
    printer.setOutputFileName(str(path))
    document.print(printer)


def export_to_html(markdown_text: str, path: str):
    document = QTextDocument()
    document.setMarkdown(markdown_text)
    write_text_file_atomic(path, document.toHtml(), "utf-8")


def _add_runs(paragraph, text: str):
    for chunk in _BOLD_ITALIC.split(text):
        if not chunk:
            continue
        if chunk.startswith("**") and chunk.endswith("**"):
            paragraph.add_run(chunk[2:-2]).bold = True
        elif chunk.startswith("*") and chunk.endswith("*"):
            paragraph.add_run(chunk[1:-1]).italic = True
        else:
            paragraph.add_run(chunk)


def export_to_docx(markdown_text: str, path: str):
    document = Document()
    for line in markdown_text.splitlines():
        stripped = line.strip()
        if not stripped:
            continue

        heading_match = _HEADING.match(stripped)
        bullet_match = _BULLET.match(stripped)
        numbered_match = _NUMBERED.match(stripped)

        if heading_match:
            level = len(heading_match.group(1))
            document.add_heading(heading_match.group(2), level=min(level, 9))
        elif bullet_match:
            paragraph = document.add_paragraph(style="List Bullet")
            _add_runs(paragraph, bullet_match.group(1))
        elif numbered_match:
            paragraph = document.add_paragraph(style="List Number")
            _add_runs(paragraph, numbered_match.group(1))
        else:
            paragraph = document.add_paragraph()
            _add_runs(paragraph, stripped)

    Path(path).parent.mkdir(parents=True, exist_ok=True)
    document.save(str(path))
